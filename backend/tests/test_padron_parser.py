"""
test_padron_parser.py — TDD suite para C-09 parser de archivos de padrón.

Cubre tasks 6.1–6.8 del tasks.md.

Ciclo: RED → GREEN → TRIANGULATE (≥2 casos) → REFACTOR.

NO usa DB — tests puramente unitarios sobre el parser.
"""
import csv
import io
import pytest

from app.schemas.padron import PadronRowDTO


# ---------------------------------------------------------------------------
# Helpers para construir archivos de prueba en memoria
# ---------------------------------------------------------------------------

def _make_xlsx_bytes(
    rows: list[dict],
    columns: list[str] | None = None,
) -> bytes:
    """Construye un xlsx en memoria con las columnas y filas dadas."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    if columns is None and rows:
        columns = list(rows[0].keys())
    elif columns is None:
        columns = []

    ws.append(columns)
    for row in rows:
        ws.append([row.get(c, "") for c in columns])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_csv_bytes(
    rows: list[dict],
    columns: list[str] | None = None,
) -> bytes:
    """Construye un csv en memoria con las columnas y filas dadas."""
    if columns is None and rows:
        columns = list(rows[0].keys())
    elif columns is None:
        columns = []

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({c: row.get(c, "") for c in columns})

    return buf.getvalue().encode("utf-8")


SAMPLE_ROWS = [
    {"nombre": "Ana", "apellidos": "García", "email": "ana@example.com", "comision": "A", "regional": "Norte"},
    {"nombre": "Bob", "apellidos": "López", "email": "bob@example.com", "comision": "B", "regional": "Sur"},
]

REQUIRED_COLUMNS = ["nombre", "apellidos", "email", "comision", "regional"]


# ---------------------------------------------------------------------------
# Task 6.1 [RED] → 6.2 [GREEN]: parse xlsx válido
# ---------------------------------------------------------------------------

def test_parse_xlsx_valid():
    """
    RED: parse_padron_file con xlsx válido retorna lista de PadronRowDTO.
    """
    from app.services.padron_parser import parse_padron_file

    xlsx_bytes = _make_xlsx_bytes(SAMPLE_ROWS)
    result = parse_padron_file(xlsx_bytes, "padron.xlsx")

    assert len(result) == 2
    assert all(isinstance(r, PadronRowDTO) for r in result)
    assert result[0].nombre == "Ana"
    assert result[0].apellidos == "García"
    assert result[0].email == "ana@example.com"
    assert result[0].comision == "A"
    assert result[0].regional == "Norte"
    assert result[1].nombre == "Bob"


# ---------------------------------------------------------------------------
# Task 6.3 [RED] → 6.4 [GREEN]: parse csv válido
# ---------------------------------------------------------------------------

def test_parse_csv_valid():
    """
    RED: parse_padron_file con csv válido retorna misma estructura PadronRowDTO.
    """
    from app.services.padron_parser import parse_padron_file

    csv_bytes = _make_csv_bytes(SAMPLE_ROWS)
    result = parse_padron_file(csv_bytes, "padron.csv")

    assert len(result) == 2
    assert all(isinstance(r, PadronRowDTO) for r in result)
    assert result[0].email == "ana@example.com"
    assert result[1].email == "bob@example.com"


# ---------------------------------------------------------------------------
# Triangulación: columnas opcionales (comision/regional) pueden estar vacías
# ---------------------------------------------------------------------------

def test_parse_xlsx_optional_columns_empty():
    """
    Triangulación: comision y regional son opcionales — pueden estar vacías.
    """
    from app.services.padron_parser import parse_padron_file

    rows = [{"nombre": "Carlos", "apellidos": "Ruiz", "email": "carlos@test.com", "comision": "", "regional": ""}]
    xlsx_bytes = _make_xlsx_bytes(rows)
    result = parse_padron_file(xlsx_bytes, "padron.xlsx")

    assert len(result) == 1
    # comision/regional vacíos → None o string vacío, ambos aceptables
    assert result[0].nombre == "Carlos"


def test_parse_csv_optional_columns_empty():
    """
    Triangulación: csv con comision y regional vacíos — no falla.
    """
    from app.services.padron_parser import parse_padron_file

    rows = [{"nombre": "Diana", "apellidos": "Pérez", "email": "diana@test.com", "comision": "", "regional": ""}]
    csv_bytes = _make_csv_bytes(rows)
    result = parse_padron_file(csv_bytes, "padron.csv")

    assert len(result) == 1
    assert result[0].nombre == "Diana"


# ---------------------------------------------------------------------------
# Task 6.5 [RED] → 6.6 [GREEN]: columna requerida faltante
# ---------------------------------------------------------------------------

def test_parse_missing_required_column():
    """
    RED: xlsx sin columna 'email' → PadronValidationError con columna en detail.
    """
    from app.services.padron_parser import parse_padron_file, PadronValidationError

    rows = [{"nombre": "Ana", "apellidos": "García", "comision": "A", "regional": "Norte"}]
    xlsx_bytes = _make_xlsx_bytes(rows, columns=["nombre", "apellidos", "comision", "regional"])

    with pytest.raises(PadronValidationError) as exc_info:
        parse_padron_file(xlsx_bytes, "padron.xlsx")

    assert exc_info.value.status_code == 422
    assert "email" in str(exc_info.value.detail).lower()


def test_parse_missing_nombre_column():
    """
    Triangulación: xlsx sin columna 'nombre' → PadronValidationError.
    """
    from app.services.padron_parser import parse_padron_file, PadronValidationError

    rows = [{"apellidos": "García", "email": "ana@test.com"}]
    xlsx_bytes = _make_xlsx_bytes(rows, columns=["apellidos", "email"])

    with pytest.raises(PadronValidationError) as exc_info:
        parse_padron_file(xlsx_bytes, "padron.xlsx")

    assert exc_info.value.status_code == 422
    assert "nombre" in str(exc_info.value.detail).lower()


# ---------------------------------------------------------------------------
# Task 6.7 [RED] → 6.8 [GREEN]: excede máximo de filas
# ---------------------------------------------------------------------------

def test_parse_exceeds_max_rows():
    """
    RED: archivo con PADRON_MAX_ROWS + 1 filas → PadronValidationError con límite.
    """
    from app.services.padron_parser import parse_padron_file, PadronValidationError
    from app.core.config import Settings

    settings = Settings()
    max_rows = settings.PADRON_MAX_ROWS

    # Construir xlsx con max_rows + 1 filas usando openpyxl directo
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(REQUIRED_COLUMNS)
    for i in range(max_rows + 1):
        ws.append([f"Nombre{i}", f"Apellido{i}", f"u{i}@test.com", "A", "Norte"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    with pytest.raises(PadronValidationError) as exc_info:
        parse_padron_file(xlsx_bytes, "padron.xlsx")

    assert exc_info.value.status_code == 422
    # El mensaje debe mencionar el límite
    detail_str = str(exc_info.value.detail).lower()
    assert str(max_rows) in detail_str or "limit" in detail_str or "máximo" in detail_str or "max" in detail_str


def test_parse_at_max_rows_succeeds():
    """
    Triangulación: exactamente PADRON_MAX_ROWS filas → no falla.
    """
    from app.services.padron_parser import parse_padron_file
    from app.core.config import Settings

    settings = Settings()
    max_rows = settings.PADRON_MAX_ROWS

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(REQUIRED_COLUMNS)
    for i in range(max_rows):
        ws.append([f"Nombre{i}", f"Apellido{i}", f"u{i}@test.com", "A", "Norte"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    result = parse_padron_file(xlsx_bytes, "padron.xlsx")
    assert len(result) == max_rows
